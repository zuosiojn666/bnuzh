#!/usr/bin/env python3
"""生成《液体表面张力系数的测定》实验报告第一至第四部分。"""

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_run_font(run, size=12, bold=False, cn="宋体", en="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:cs"), en)
    rFonts.set(qn("w:eastAsia"), cn)


def set_paragraph_format(paragraph, before=0, after=0, line=1.5, indent=True, align="left"):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if indent:
        fmt.first_line_indent = Pt(24)
    else:
        fmt.first_line_indent = Pt(0)
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


_SUB = re.compile(r"_\{([^}]+)\}")


def write_marked(paragraph, text, size=12, bold=False, cn="宋体"):
    """把 F_{T}、U_{max} 这类标记写成 Word 下标。"""
    pos = 0
    for match in _SUB.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, bold=bold, cn=cn)
        sub = paragraph.add_run(match.group(1))
        set_run_font(sub, size=size, bold=bold, cn=cn)
        sub.font.subscript = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=bold, cn=cn)


def add_text(doc, text, size=12, bold=False, cn="宋体", before=0, after=0, indent=True, align="left"):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, indent=indent, align=align)
    write_marked(p, text, size=size, bold=bold, cn=cn)
    return p


def add_mixed(doc, parts, before=0, after=0, indent=True, align="left", line=1.5):
    """parts: list of (text, size, bold, cn)."""
    p = doc.add_paragraph()
    set_paragraph_format(p, before=before, after=after, indent=indent, align=align, line=line)
    for text, size, bold, cn in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, cn=cn)
    return p


def add_formula(doc, expr, number):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=3, after=3, indent=False, align="center")
    write_marked(p, expr + "      " + number, size=12)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "宋体")

    add_text(doc, "北京师范大学文理学院物理系", size=16, bold=True, cn="黑体", indent=False, align="center", before=18, after=6)
    add_text(doc, "物理实验报告", size=22, bold=True, cn="黑体", indent=False, align="center", before=6, after=18)

    cover_lines = [
        "课程名称：大学物理实验 B",
        "实验项目：液体表面张力系数的测定",
        "学生年级：                    专业：",
        "学生姓名：                    学号：",
        "指导教师：                    实验教师：",
        "实验日期：        年      月      日",
    ]
    for line in cover_lines:
        add_text(doc, line, size=14, indent=False, align="center", before=6, after=6)

    add_text(
        doc,
        "第一至第四部分",
        size=12,
        indent=False,
        align="center",
        before=18,
        after=6,
    )

    # ---------- 一 ----------
    add_text(doc, "一、实验目的", size=14, bold=True, cn="黑体", indent=False, before=12, after=4)
    for item in (
        "1. 掌握用拉脱法测量液体表面张力系数的原理和方法。",
        "2. 掌握硅压阻力敏传感器的定标方法。",
    ):
        add_text(doc, item, before=1, after=1)

    # ---------- 二 ----------
    add_text(doc, "二、实验原理", size=14, bold=True, cn="黑体", indent=False, before=10, after=4)
    add_text(doc, "1. 原理简述", size=12, bold=True, cn="黑体", indent=False, before=4, after=2)
    add_text(doc, "液体内部分子受力抵消。表面分子所受合力指向液体内部，液面趋于收缩，此力即表面张力。α 为表面张力系数，单位 N/m。")
    add_text(doc, "拉脱法通过金属环脱离液面所需的力求 α。液膜同时附在环的内外沿，接触线长为 π(D₁+D₂)，故")
    add_formula(doc, "F = α π (D₁ + D₂)", "(1)")
    add_text(doc, "拉断前，环受拉力 F_{T}、重力 F_{G}、表面张力 F 和液膜重力 F_{W}：")
    add_formula(doc, "F_{T} = F_{G} + F cosθ + F_{W}", "(2)")
    add_formula(doc, "F_{W} = π (D₁² − D₂²) h ρ g / 4", "(3)")
    add_text(doc, "θ 为 F 与竖直方向的夹角。h 为液膜拉出高度，h<0 时该项为浮力。")
    add_text(doc, "硅压阻力敏传感器由弹性梁和四个力敏电阻组成的电桥构成，输出电压与拉力成正比：")
    add_formula(doc, "U = K F_{T}", "(4)")
    add_text(doc, "K 的单位为 mV/N。拉力最大时 θ=0，电压为 U_{max}；拉断后电压为 U₂，此时只剩环重。于是")
    add_formula(doc, "F_{T,max} − F₂ = (U_{max} − U₂) / K = F + F_{W}", "(5)")
    add_text(doc, "联立得")
    add_formula(doc, "α = [(U_{max}−U₂)/K − π(D₁²−D₂²)h_{max}ρg/4] / [π(D₁+D₂)]", "(6)")
    add_text(doc, "讲义指出液膜重力可以忽略，实验采用")
    add_formula(doc, "α = (U_{max} − U₂) / [π K (D₁ + D₂)]", "(7)")
    add_text(doc, "定标时 F=mg，取 g=9.8 m/s²。对 U–F 作直线拟合，斜率即为 K：")
    add_formula(doc, "U = K F + b", "(8)")

    add_text(doc, "2. 预习思考题", size=12, bold=True, cn="黑体", indent=False, before=8, after=2)
    add_text(doc, "（1）液体表面张力产生的原因", bold=True, before=4, after=1)
    add_text(doc, "内部分子各向受力抵消。表面分子外侧分子稀少，合力指向液内，液面收缩，形成表面张力。")
    add_text(doc, "（2）表面张力系数与哪些因素有关", bold=True, before=4, after=1)
    add_text(doc, "主要与液体种类、温度和杂质（浓度）有关。温度升高，α 减小；肥皂等表面活性物质使 α 明显下降。α 与吊环尺寸无关。")
    add_text(doc, "（3）吊环不水平的影响", bold=True, before=4, after=1)
    add_text(doc, "液膜不能整圈同时拉脱，有效周长小于 π(D₁+D₂)，且最大拉力处 θ 不全为 0。U_{max} 偏小，测得的 α 偏小，读数也不稳定。")
    add_text(doc, "（4）拉膜过程中电压如何变化", bold=True, before=4, after=1)
    add_text(doc, "浸入时浮力使电压下降；液面下降后浮力减小，电压回升。形成液膜后，表面张力的竖直分力增大，电压升至 U_{max}。液膜破裂后骤降至 U₂。")

    # ---------- 三 ----------
    add_text(doc, "三、实验仪器", size=14, bold=True, cn="黑体", indent=False, before=10, after=4)
    add_text(doc, "DH4607J 型液体表面张力系数测定仪及测试架（压力传感器、吊环、加热器皿、升降台、光电门）、0.5 g 砝码 7 个、砝码盘、镊子、水平泡、游标卡尺、超声波清洗机、蒸馏水、肥皂水。")

    # ---------- 四 ----------
    add_text(doc, "四、实验步骤", size=14, bold=True, cn="黑体", indent=False, before=10, after=4)
    steps = [
        "1. 预热不少于 10 分钟，调节底座使水平泡居中。",
        "2. 用游标卡尺测量吊环外径 D₁、内径 D₂。超声波清洗吊环后晾干，勿划伤或折弯。",
        "3. 进入灵敏度测量界面，挂上砝码盘并置零。用镊子依次加 0.5 g 至 3.5 g 砝码，稳定后录入电压。取 g=9.8 m/s²，由 U–F 直线斜率得到 K。",
        "4. 进入表面张力测量界面，挂环并置零，阈值设为 0.01 mV。注入蒸馏水，使环的下沿浸没，位移置零。缓慢下降液面至液膜破裂，读出 U_{max} 和 U₂。重复 6 次。",
        "5. 将水温依次设为 30 ℃、40 ℃、50 ℃、60 ℃，温度稳定后各测一组。",
        "6. 洗净吊环和器皿，换肥皂水，室温下测 3 次。",
        "7. 砝码用镊子夹取；拉膜时吊环勿晃动，液面下降宜慢；蒸馏水保持洁净。",
    ]
    for item in steps:
        add_text(doc, item, before=1, after=1)

    out = "/workspace/docs/液体表面张力系数的测定_实验报告_第一至四部分.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
